"""Exports generated from actual case records: PDF, JSON, Markdown.

PDF uses ReportLab with clear headings and page numbers. JSON export includes
structured findings, evidence IDs, relationships, timeline events, research
provenance and review decisions.
"""
from __future__ import annotations

import io
import json
from datetime import datetime

from ..db import repo
from ..schemas import CaseWorkspace
from .report import DISCLAIMER


def workspace(case_id: str) -> CaseWorkspace:
    case = repo.get("cases", case_id)
    return CaseWorkspace(
        case=case,
        evidence=repo.list("evidence", case_id),
        observations=repo.list("observations", case_id),
        entities=repo.list("entities", case_id),
        relationships=repo.list("relationships", case_id),
        timeline=repo.list("timeline_events", case_id),
        findings=repo.list("findings", case_id),
        research_queries=repo.list("research_queries", case_id),
        research_results=repo.list("research_results", case_id),
    )


def to_json(case_id: str, report_id: str | None = None) -> bytes:
    ws = workspace(case_id)
    payload = ws.model_dump(mode="json")
    payload["review_actions"] = [r.model_dump(mode="json") for r in repo.list("review_actions", case_id)]
    payload["audit_events"] = [a.model_dump(mode="json") for a in repo.list("audit_events", case_id)]
    if report_id:
        rep = repo.get("report_versions", report_id)
        if rep:
            payload["report"] = rep.model_dump(mode="json")
    payload["_meta"] = {
        "disclaimer": DISCLAIMER,
        "exported_at": datetime.utcnow().isoformat() + "Z",
        "note": "Synthetic demonstration data. Not a certified forensic export.",
    }
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")


def _pdf(text: str, title: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.enums import TA_CENTER

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm,
                            leftMargin=18 * mm, rightMargin=18 * mm, title=title)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=15, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9.5, leading=13)
    small = ParagraphStyle("small", parent=body, textColor="#666666")
    center = ParagraphStyle("center", parent=small, alignment=TA_CENTER)

    story = []

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColorRGB(0.4, 0.4, 0.4)
        canvas.drawString(18 * mm, 10 * mm, DISCLAIMER)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            story.append(Spacer(1, 3))
            continue
        safe = (line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
        if line.startswith("# "):
            story.append(Paragraph(safe[2:], h1))
        elif line.startswith("## "):
            story.append(Paragraph(safe[3:], h2))
        elif line.startswith("**Demonstration"):
            story.append(Paragraph(safe.replace("**", ""), center))
        elif safe.startswith("- "):
            story.append(Paragraph("&bull; " + safe[2:], body))
        else:
            story.append(Paragraph(safe.replace("**", "<b>").replace("`", ""), body))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def to_pdf(case_id: str, report_id: str) -> bytes:
    rep = repo.get("report_versions", report_id)
    title = repo.get("cases", case_id).title
    return _pdf(rep.body_markdown, title)


def to_markdown(case_id: str, report_id: str) -> bytes:
    return repo.get("report_versions", report_id).body_markdown.encode("utf-8")


def to_docx(case_id: str, report_id: str) -> bytes | None:
    """DOCX only if python-docx is present; otherwise None (feature omitted honestly)."""
    try:
        import docx  # noqa: F401
    except Exception:  # noqa: BLE001
        return None
    from docx import Document
    rep = repo.get("report_versions", report_id)
    doc = Document()
    for line in rep.body_markdown.splitlines():
        s = line.strip()
        if s.startswith("# "):
            doc.add_heading(s[2:], 0)
        elif s.startswith("## "):
            doc.add_heading(s[3:], 1)
        elif s.startswith("- "):
            doc.add_paragraph(s[2:], style="List Bullet")
        elif s:
            doc.add_paragraph(s)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
